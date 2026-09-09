from pathlib import Path

from openpyxl import load_workbook

from sas_lineage.exporter import (
    COLUMNS,
    edge_rows,
    export_excel,
    issue_rows,
    make_fileid,
    schema_of,
)
from sas_lineage.pipeline import analyze_source

EXAMPLES = Path(__file__).resolve().parent.parent.parent / "examples"

SIMPLE = "libname raw '/data/raw';\ndata work.x;\n set raw.y;\nrun;"


def test_schema_of():
    assert schema_of("raw.customer") == "raw"
    assert schema_of("WORK.X") == "work"
    assert schema_of("plain") == "work"  # unqualified defaults to work


def test_column_schema_is_exact():
    assert COLUMNS == [
        "fileid", "block_id", "source_ref_id", "target_ref_id",
        "source_canonical_name", "target_canonical_name",
        "source_db_schema", "target_db_schema",
        "source_lib_path", "target_ref_path",
        "dst_source_db_schema", "dst_target_db_schema",
        "log_verified",
    ]


def test_edge_row_contents(plain):
    rows = edge_rows(analyze_source(SIMPLE), fileid="f1")
    assert len(rows) == 1
    row = rows[0]
    assert row["fileid"] == "f1"
    assert plain(row["block_id"]) == "b_1"
    assert plain(row["source_ref_id"]) == "b_1:t_1"
    assert plain(row["target_ref_id"]) == "b_1:t_2"
    assert row["source_canonical_name"] == "raw.y"
    assert row["target_canonical_name"] == "work.x"
    assert row["source_db_schema"] == "raw"
    assert row["target_db_schema"] == "work"
    assert row["source_lib_path"] == "/data/raw"  # from the libname statement
    assert row["target_ref_path"] == ""  # work has no libname
    assert row["log_verified"] == ""


def test_dst_schemas_from_schema_map():
    rows = edge_rows(analyze_source(SIMPLE), "f1",
                     schema_map={"raw": "landing", "work": "staging"})
    assert rows[0]["dst_source_db_schema"] == "landing"
    assert rows[0]["dst_target_db_schema"] == "staging"


def test_dst_schemas_empty_by_default():
    rows = edge_rows(analyze_source(SIMPLE), "f1")
    assert rows[0]["dst_source_db_schema"] == ""
    assert rows[0]["dst_target_db_schema"] == ""


def test_file_flow_block_id_is_reader_block(plain):
    src = "data work.a;\n set raw.x;\nrun;\ndata work.b;\n set work.a;\nrun;"
    rows = edge_rows(analyze_source(src), "f1")
    file_flow = [r for r in rows
                 if plain(r["source_ref_id"]).startswith("b_1")
                 and plain(r["target_ref_id"]).startswith("b_2")]
    assert len(file_flow) == 1
    # the block where the read lands
    assert plain(file_flow[0]["block_id"]) == "b_2"


def test_issue_rows_capture_non_parsed_blocks():
    analysis = analyze_source("data work.x;\n set &nolib..y;\nrun;")
    issues = issue_rows(analysis, "f1")
    assert len(issues) == 1
    assert issues[0]["status"] == "PARTIAL"
    assert "&nolib" in issues[0]["unresolved"]


def test_fileid_is_stable_relative_path(tmp_path):
    p = tmp_path / "proj" / "a.sas"
    p.parent.mkdir()
    p.write_text("data w.x; run;")
    fid1 = make_fileid(p, tmp_path)
    p.write_text("data w.y; run;")           # content change
    assert make_fileid(p, tmp_path) == fid1 == "proj/a.sas"


def test_fileid_carries_no_content_hash(tmp_path):
    f = tmp_path / "job.sas"
    f.write_text("data a; set b; run;")
    assert "#" not in make_fileid(f, tmp_path)
    assert make_fileid(f, tmp_path) == f"{tmp_path.name}/job.sas"


def test_fileid_without_base_dir_keeps_the_parent_folder(tmp_path):
    f = tmp_path / "etl" / "job.sas"
    f.parent.mkdir()
    f.write_text("data a; set b; run;")
    assert make_fileid(f) == "etl/job.sas"


def test_export_excel_end_to_end(tmp_path):
    out = export_excel([EXAMPLES / "ex9_libname_pipeline.sas"],
                       tmp_path / "lineage.xlsx")
    wb = load_workbook(out)
    assert wb.sheetnames == ["lineage", "issues"]
    sheet = wb["lineage"]
    header = [c.value for c in sheet[1]]
    assert header == COLUMNS
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    # 3 BLOCK_FLOW + 2 FILE_FLOW edges
    assert len(rows) == 5
    by_col = [dict(zip(COLUMNS, r)) for r in rows]
    final = [r for r in by_col if r["target_canonical_name"] == "mart.orders_final"]
    assert final[0]["target_ref_path"] == "/data/warehouse/mart"
    first = [r for r in by_col if r["source_canonical_name"] == "raw.orders"]
    assert first[0]["source_lib_path"] == "/data/landing/raw"
    assert all(r["log_verified"] in (None, "") for r in by_col)


def test_export_excel_multiple_files_distinct_fileids(tmp_path):
    out = export_excel(
        [EXAMPLES / "ex2_merge_in_flags.sas", EXAMPLES / "ex9_libname_pipeline.sas"],
        tmp_path / "multi.xlsx",
    )
    sheet = load_workbook(out)["lineage"]
    fileids = {r[0] for r in sheet.iter_rows(min_row=2, values_only=True)}
    assert len(fileids) == 2


def test_export_excel_passes_injected_block_ids(tmp_path):
    out = export_excel([EXAMPLES / "ex10_annotated_blocks.sas"],
                       tmp_path / "annotated.xlsx")
    sheet = load_workbook(out)["lineage"]
    rows = [dict(zip(COLUMNS, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    assert {r["block_id"] for r in rows} == {
        "1:2645661869534515", "2:7118293441055627",
    }
    sql_rows = [r for r in rows if r["block_id"] == "1:2645661869534515"]
    assert sql_rows[0]["source_ref_id"] == "1:2645661869534515:t_1"
    assert all(r["fileid"] == "examples/ex10_annotated_blocks.sas" for r in rows)


def test_export_excel_macro_instance_block_ids(tmp_path):
    out = export_excel([EXAMPLES / "ex11_annotated_macro.sas"],
                       tmp_path / "macro.xlsx")
    sheet = load_workbook(out)["lineage"]
    rows = [dict(zip(COLUMNS, r)) for r in sheet.iter_rows(min_row=2, values_only=True)]
    instance_rows = {r["block_id"] for r in rows if "#" in r["block_id"]}
    # within-macro edges carry the macro DEFINITION's injected id + instance
    assert instance_rows == {"1:5550001112223334#1", "1:5550001112223334#2"}
    file_edge = [r for r in rows if r["source_ref_id"].startswith("1:5550001112223334#1")
                 and r["target_ref_id"].startswith("4:")]
    assert file_edge, "FILE_FLOW row linking macro instance to caller block missing"


def test_export_excel_issues_sheet(tmp_path):
    bad = tmp_path / "bad.sas"
    bad.write_text("data work.x;\n set &nope..y;\nrun;")
    out = export_excel([bad], tmp_path / "issues.xlsx")
    issues = load_workbook(out)["issues"]
    rows = list(issues.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][2] == "PARTIAL"
