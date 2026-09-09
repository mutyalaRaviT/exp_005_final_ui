import io
import threading
import time
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import create_app

# Same shape as Task 4's fixture: seed -> mid -> leaf -> seed is a loop
# (work.t3 feeds back into the seed); report hangs off mid, outside the loop.
FIXTURE = {
    "01_seed.sas": "data work.t1; set work.t3; run;",
    "02_mid.sas": "data work.t2; set work.t1; run;",
    "03_leaf.sas": "data work.t3; set work.t2; run;",
    "04_report.sas": "data work.t4; set work.t2; run;",
}


def make_repo(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    for name, code in FIXTURE.items():
        (root / name).write_text(code)
    return root


def _wait_ready(client, timeout=10):
    deadline = time.time() + timeout
    while time.time() < deadline:
        if client.get("/api/status").json()["state"] == "ready":
            return
        time.sleep(0.02)
    raise TimeoutError("index never became ready")


@pytest.fixture
def client(tmp_path):
    root = make_repo(tmp_path)
    app = create_app([root], tmp_path / "db.duckdb")
    with TestClient(app) as c:
        _wait_ready(c)
        yield c


def _seed_fileid(client):
    hits = client.get("/api/search", params={"q": "01_seed"}).json()["hits"]
    (hit,) = [h for h in hits if h["kind"] == "file"]
    (fileid,) = hit["files"]
    return fileid


def test_files_lists_every_indexed_file_before_anything_is_parsed(client):
    """The explorer's ALL FILES layer is fed by this endpoint on the first
    page load — before any neighborhood has caused a file to be parsed. So it
    has to answer with the WHOLE index, not only the files already ingested.
    """
    files = client.get("/api/files").json()["files"]
    assert {f["label"] for f in files} == set(FIXTURE)
    assert all(f["id"] and f["folder"] for f in files)
    # deterministic order, so the same index always draws the same tree
    assert [f["id"] for f in files] == sorted(f["id"] for f in files)


def test_files_still_lists_a_file_after_it_is_parsed(client):
    """Once a file IS ingested it must appear exactly once, not twice."""
    fileid = _seed_fileid(client)
    client.get(f"/api/file/{quote(fileid, safe='')}")
    files = client.get("/api/files").json()["files"]
    assert [f["id"] for f in files].count(fileid) == 1
    assert len(files) == len(FIXTURE)


def test_files_and_edges_can_be_asked_for_at_the_same_time(client):
    """The screen fires several requests at once on its very first paint
    (status + files + the grid's whole-view Arrow load), and they all land on
    ONE DuckDB connection. A smoke check that the pair coexists; the
    guarantee itself is asserted deterministically in
    test_service.py::test_list_files_is_serialized.
    """
    errors = []

    def hammer_files():
        for _ in range(10):
            response = client.get("/api/files")
            if response.status_code != 200:
                errors.append(response.status_code)
                return
            if any(not row["label"] for row in response.json()["files"]):
                errors.append("a file row came back without its name")
                return

    def hammer_edges():
        for _ in range(10):
            client.get("/api/edges/arrow", params={"view": "final"})

    threads = [threading.Thread(target=hammer_files),
               threading.Thread(target=hammer_edges)]
    for t in threads:
        t.start()
    for t in threads:
        t.join(timeout=60)

    assert errors == []


def test_status_shape(client):
    body = client.get("/api/status").json()
    assert body["state"] == "ready"
    assert body["total"] == 4
    assert body["scanner"] in ("rg", "python")
    assert body["tables_known"] > 0


def test_search_returns_hits(client):
    body = client.get("/api/search", params={"q": "work.t1"}).json()
    assert any(h["kind"] == "table" and h["value"] == "work.t1"
               for h in body["hits"])


def test_neighborhood_shape(client):
    fileid = _seed_fileid(client)
    body = client.get("/api/neighborhood",
                       params={"file": fileid, "depth": 2}).json()
    assert {n["label"] for n in body["nodes"]} == set(FIXTURE)
    assert body["edges"]
    assert all(e["level"] == "project" and e["provenance"] == "inferred"
               for e in body["edges"])
    assert any("->" in line for line in body["story"])
    assert fileid in body["order"]


def test_neighborhood_takes_directional_depths(client):
    """`up=`/`down=` bound the walk per direction; nodes carry `role` and
    the payload names its `seeds` so the graph can ring/tint them."""
    fileid = _seed_fileid(client)
    body = client.get("/api/neighborhood",
                      params={"file": fileid, "up": 0, "down": 1}).json()
    assert body["seeds"] == [fileid]
    roles = {n["label"]: n["role"] for n in body["nodes"]}
    assert roles["01_seed.sas"] == "seed"
    assert all(r in ("seed", "down") for r in roles.values())


def test_file_detail_shape(client):
    fileid = _seed_fileid(client)
    body = client.get(f"/api/file/{fileid}").json()
    assert body["name"] == "01_seed.sas"
    assert body["block_edges"]
    assert {e["provenance"] for e in body["block_edges"]} == {"fact"}


def test_file_detail_accepts_a_percent_encoded_fileid(client):
    # a fileid is a path now, and the web client sends it through
    # encodeURIComponent -> "repo%2F01_seed.sas"
    fileid = _seed_fileid(client)
    assert "/" in fileid
    body = client.get(f"/api/file/{quote(fileid, safe='')}").json()
    assert body["fileid"] == fileid
    assert body["name"] == "01_seed.sas"


def test_order_defaults_to_all_indexed(client):
    body = client.get("/api/order").json()
    assert len(body) == 4
    assert all("score" in row and "cyclic" in row for row in body.values())


def test_order_with_explicit_files(client):
    fileid = _seed_fileid(client)
    body = client.get("/api/order", params={"files": fileid}).json()
    assert list(body) == [fileid]
    assert body[fileid]["score"] == 0
    assert body[fileid]["cyclic"] is False


def test_excel_download_shape(client):
    fileid = _seed_fileid(client)
    resp = client.get("/api/excel", params={"files": fileid})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    disposition = resp.headers["content-disposition"]
    assert disposition.startswith("attachment; filename=lineage_")
    assert disposition.endswith(".xlsx")
    book = load_workbook(io.BytesIO(resp.content))
    assert book.sheetnames == ["provided", "customer_changes", "final_edges"]


def test_post_edit_creates_edit_id(client):
    resp = client.post("/api/edits", json={
        "action": "confirm", "src": "a.sas", "dst": "b.sas",
        "table_name": "work.t1", "level": "project", "editor": "ravi",
        "comment": "ok"})
    assert resp.status_code == 200
    assert resp.json()["edit_id"]


def test_post_edit_bad_action_is_422(client):
    resp = client.post("/api/edits", json={
        "action": "delete", "src": "a.sas", "dst": "b.sas",
        "table_name": "work.t1", "level": "project", "editor": "ravi",
        "comment": ""})
    assert resp.status_code == 422


def test_cors_allows_the_dev_frontend(client):
    resp = client.get("/api/status",
                       headers={"Origin": "http://localhost:5173"})
    assert (resp.headers.get("access-control-allow-origin")
            == "http://localhost:5173")


# ---- Task 10: /api/edges ----------------------------------------------------

def test_edges_shape(client):
    body = client.get("/api/edges").json()
    assert "total" in body and "rows" in body
    assert body["total"] > 0
    assert all(r["src"] and r["dst"] and r["level"] and r["provenance"]
               for r in body["rows"])


def test_edges_filter_and_limit(client):
    body = client.get("/api/edges", params={"filter": "t1", "limit": 1}).json()
    assert body["total"] >= 1
    assert len(body["rows"]) == 1


def test_edges_level_filter(client):
    body = client.get("/api/edges", params={"level": "block"}).json()
    assert body["rows"]
    assert all(r["level"] == "block" for r in body["rows"])


def test_edges_paging_offset_beyond_total(client):
    total = client.get("/api/edges").json()["total"]
    body = client.get("/api/edges", params={"offset": total + 100}).json()
    assert body["rows"] == []
    assert body["total"] == total


# ---- Task 14: GET /api/edges/arrow -------------------------------------------

def _arrow_table(resp):
    import pyarrow as pa

    return pa.ipc.open_stream(pa.BufferReader(resp.content)).read_all()


def test_edges_arrow_content_type_and_columns(client):
    resp = client.get("/api/edges/arrow")
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.apache.arrow.stream"
    table = _arrow_table(resp)
    assert table.schema.names == ["fileid", "src", "dst", "table_name",
                                  "level", "provenance", "block_id",
                                  "requires_check", "freshness"]


def test_edges_arrow_round_trips_the_same_row_count_as_the_json_view(client):
    total = client.get("/api/edges").json()["total"]
    assert _arrow_table(client.get("/api/edges/arrow")).num_rows == total


def test_edges_arrow_engine_view_has_no_human_gold(client):
    fileid = _seed_fileid(client)
    client.post("/api/edits", json={
        "action": "confirm", "src": fileid, "dst": "work.t1",
        "table_name": "work.t3", "level": "block", "editor": "ravi",
        "comment": "ok"})
    engine = _arrow_table(
        client.get("/api/edges/arrow", params={"view": "engine"})).to_pydict()
    assert engine["provenance"]
    assert set(engine["provenance"]) <= {"fact", "inferred"}


def test_edges_arrow_human_view_returns_the_customer_edits(client):
    client.post("/api/edits", json={
        "action": "add", "src": "01_seed.sas", "dst": "05_extra.sas",
        "table_name": "work.extra", "level": "project", "editor": "ravi",
        "comment": "new flow"})
    human = _arrow_table(
        client.get("/api/edges/arrow", params={"view": "human"})).to_pydict()
    assert "05_extra.sas" in human["dst"]
    assert set(human["provenance"]) == {"human_gold"}


def test_edges_arrow_bad_view_is_422(client):
    resp = client.get("/api/edges/arrow", params={"view": "excel"})
    assert resp.status_code == 422


def test_post_edit_carries_the_block_it_was_made_against(client):
    """A confirm from the grid pins itself to that row's block id, so the
    final view can upgrade exactly that block's flow to human_gold."""
    row = next(r for r in client.get(
        "/api/edges", params={"level": "block"}).json()["rows"] if r["block_id"])
    resp = client.post("/api/edits", json={
        "action": "confirm", "src": row["src"], "dst": row["dst"],
        "table_name": row["tables"][0], "level": "block", "editor": "ravi",
        "comment": "", "block_id": row["block_id"]})
    assert resp.status_code == 200

    after = client.get("/api/edges", params={"level": "block"}).json()["rows"]
    upgraded = [r for r in after if r["block_id"] == row["block_id"]]
    assert upgraded
    assert all(r["provenance"] == "human_gold" for r in upgraded)


def test_post_edit_correct_with_a_fileid_round_trips(client):
    """The edit form sends the file the edge belongs to: a block id is only
    unique WITHIN a file, so a block-level edit has to name both."""
    fileid = _seed_fileid(client)
    row = next(r for r in client.get(
        "/api/edges", params={"level": "block"}).json()["rows"] if r["block_id"])
    resp = client.post("/api/edits", json={
        "action": "correct", "src": row["src"], "dst": row["dst"],
        "table_name": "work.right", "level": "block", "editor": "ravi",
        "comment": "reads work.right", "block_id": row["block_id"],
        "fileid": fileid})
    assert resp.status_code == 200
    assert resp.json()["edit_id"]

    human = _arrow_table(
        client.get("/api/edges/arrow", params={"view": "human"})).to_pydict()
    assert human["fileid"] == [fileid]
    assert human["table_name"] == ["work.right"]
    assert human["provenance"] == ["human_gold"]

    # and the corrected flow is in the final view as human_gold
    final = _arrow_table(
        client.get("/api/edges/arrow", params={"view": "final"})).to_pydict()
    corrected = [i for i, t in enumerate(final["table_name"])
                 if t == "work.right"]
    assert corrected
    assert final["provenance"][corrected[0]] == "human_gold"
    assert final["fileid"][corrected[0]] == fileid


def test_concurrent_requests_do_not_collide_on_the_store(client):
    """The screen fires several calls at once — the grid's whole-view Arrow
    load lands while the graph, run order and paged edges are still in
    flight. They all share one DuckDB connection (and its temp views), so
    parse-on-demand has to serialize instead of ingesting the same file
    twice at once."""
    import threading

    codes: list[int] = []
    lock = threading.Lock()

    def hit(path, params):
        status = client.get(path, params=params).status_code
        with lock:
            codes.append(status)

    calls = [
        ("/api/edges/arrow", {"view": "final"}),
        ("/api/edges", {}),
        ("/api/edges/arrow", {"view": "engine"}),
        ("/api/excel", {}),
    ]
    threads = [threading.Thread(target=hit, args=call) for call in calls]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join()

    assert codes == [200] * len(calls)


# ---- T1: the review-queue endpoint -----------------------------------------

def test_pending_endpoint_lists_flagged_edits_and_empties_after_confirm(tmp_path):
    root = make_repo(tmp_path)
    app = create_app([root], tmp_path / "db.duckdb")
    with TestClient(app) as client:
        _wait_ready(client)
        assert client.get("/api/edits/pending").json() == {"pending": []}

        fileid = "repo/01_seed.sas"
        detail = client.get(f"/api/file/{quote(fileid, safe='')}").json()
        block_id = detail["blocks"][0]["id"]
        client.post("/api/edits", json={
            "action": "confirm", "src": fileid, "dst": "work.t1",
            "table_name": "work.t3", "level": "block", "editor": "ravi",
            "comment": "", "block_id": block_id, "fileid": fileid})

        (root / "01_seed.sas").write_text(
            "data work.t1; set work.t3; label t1='moved'; run;")
        client.get(f"/api/file/{quote(fileid, safe='')}")   # re-ingest

        pending = client.get("/api/edits/pending").json()["pending"]
        assert [p["fileid"] for p in pending] == [fileid]
        assert pending[0]["action"] == "confirm"

        client.post("/api/edits", json={
            "action": "confirm", "src": fileid, "dst": "work.t1",
            "table_name": "work.t3", "level": "block", "editor": "ravi",
            "comment": "still holds"})
        assert client.get("/api/edits/pending").json() == {"pending": []}


# ---- code save + timeline endpoints ----------------------------------------

def test_save_endpoint_writes_code_and_returns_delta_and_timeline(tmp_path):
    root = make_repo(tmp_path)
    app = create_app([root], tmp_path / "db.duckdb")
    with TestClient(app) as client:
        _wait_ready(client)
        fileid = "repo/01_seed.sas"
        client.get(f"/api/file/{quote(fileid, safe='')}")   # ingest

        res = client.post(f"/api/file/{quote(fileid, safe='')}/save", json={
            "code": "data work.t1; set work.t9; run;",
            "editor": "ravi", "reason": "switch seed to t9"})
        assert res.status_code == 200
        body = res.json()
        assert (root / "01_seed.sas").read_text() == \
            "data work.t1; set work.t9; run;"
        assert {"src": "work.t9", "dst": "work.t1", "table_name": "work.t9",
                "level": "block"} in body["added"]
        assert body["commit"]

        hist = client.get(
            f"/api/file/{quote(fileid, safe='')}/history").json()["history"]
        assert [h["editor"] for h in hist] == ["ravi", "engine"]
        assert hist[0]["reason"] == "switch seed to t9"

        at = client.get(
            f"/api/file/{quote(fileid, safe='')}/at/{hist[1]['commit']}").json()
        assert "set work.t3" in at["code"]

        # the code pane, asked again, serves the saved bytes
        detail = client.get(f"/api/file/{quote(fileid, safe='')}").json()
        assert "set work.t9" in detail["code"]


def test_save_endpoint_rejects_missing_reason_or_editor(tmp_path):
    root = make_repo(tmp_path)
    app = create_app([root], tmp_path / "db.duckdb")
    with TestClient(app) as client:
        _wait_ready(client)
        fileid = "repo/01_seed.sas"
        client.get(f"/api/file/{quote(fileid, safe='')}")
        res = client.post(f"/api/file/{quote(fileid, safe='')}/save", json={
            "code": "data a; run;", "editor": " ", "reason": "x"})
        assert res.status_code == 422
        res = client.post(f"/api/file/{quote(fileid, safe='')}/save", json={
            "code": "data a; run;", "editor": "ravi", "reason": ""})
        assert res.status_code == 422


def test_status_reports_sync_fields(client):
    payload = client.get("/api/status").json()
    assert "synced_at" in payload
    assert "files_on_disk" in payload
    assert "sources_indexed" in payload


def test_post_sync_runs_and_reports(client):
    payload = client.post("/api/sync").json()
    assert payload["files_on_disk"] >= 1
    assert "synced_at" in payload


def test_blocklinks_endpoint_scopes_to_files_param(tmp_path):
    root = make_repo(tmp_path)
    app = create_app([root], tmp_path / "db.duckdb")
    with TestClient(app) as client:
        _wait_ready(client)
        links = client.get("/api/blocklinks?files=repo/01_seed.sas").json()["links"]
        assert isinstance(links, list)          # single file -> no cross-file links
        assert links == []
