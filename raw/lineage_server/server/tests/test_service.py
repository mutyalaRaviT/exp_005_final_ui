import io
import threading

import pytest
from openpyxl import load_workbook

import service
from indexer import Indexer
from sas_lineage.store import LineageStore

# seed -> mid -> leaf -> seed is a loop (work.t3 feeds back into the seed);
# report hangs off mid and is the only file outside the loop.
FIXTURE = {
    "01_seed.sas": "data work.t1; set work.t3; run;",
    "02_mid.sas": "data work.t2; set work.t1; run;",
    "03_leaf.sas": "data work.t3; set work.t2; run;",
    "04_report.sas": "data work.t4; set work.t2; run;",
}


def make_repo(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    for name, code in FIXTURE.items():
        (root / name).write_text(code)
    return root


@pytest.fixture
def svc(tmp_path):
    root = make_repo(tmp_path)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    return service.Service(store, ix, [root])


def seed_fileid(svc):
    (fileid,) = [f for f in svc.indexer.files_mentioning("work.t1")
                 if f.endswith("01_seed.sas")]
    return fileid


def test_fileid_is_the_relative_path(svc):
    assert seed_fileid(svc) == "repo/01_seed.sas"


def test_list_files_lists_every_indexed_file(svc):
    """What the explorer's ALL FILES layer draws."""
    rows = svc.list_files()
    assert [r["label"] for r in rows] == sorted(FIXTURE)
    assert all(r["id"].endswith(r["label"]) and r["folder"] == "repo"
               for r in rows)


def test_indexer_and_service_share_one_lock_for_one_connection(tmp_path):
    """The Indexer and the Service hold the SAME DuckDB connection, and both
    read it while requests are in flight — the indexer answers the typeahead
    (`suggest`) and the status bar (`_tables_known`) while the grid's Arrow
    load runs through the service. Two independent locks would not stop them
    colliding on that one connection, so there is exactly one.
    """
    root = make_repo(tmp_path)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    assert svc._lock is ix.lock


@pytest.mark.parametrize("call", [
    lambda ix: ix.suggest("work"),
    lambda ix: ix.files_mentioning("work.t1"),
    lambda ix: ix._tables_known(),
])
def test_indexer_store_reads_are_serialized(svc, call):
    """Each of these answers an endpoint the screen polls (search, status),
    so each has to wait for whatever else is using the connection."""
    started, finished = threading.Event(), threading.Event()

    def run():
        started.set()
        call(svc.indexer)
        finished.set()

    with svc._lock:
        worker = threading.Thread(target=run)
        worker.start()
        assert started.wait(timeout=5)
        assert not finished.wait(timeout=0.3)

    worker.join(timeout=5)
    assert finished.is_set()


def test_list_files_is_serialized(svc):
    """One Service owns ONE DuckDB connection, so every store-touching call
    has to take the same lock — a read that skips it gets ANOTHER query's
    rows off that connection. This matters most for `list_files`: the
    explorer asks for it on the first paint, while the grid's whole-view
    Arrow load is already in flight on the same connection.
    """
    started, finished = threading.Event(), threading.Event()

    def call():
        started.set()
        svc.list_files()
        finished.set()

    with svc._lock:
        worker = threading.Thread(target=call)
        worker.start()
        assert started.wait(timeout=5)
        # the lock is held here, so the call cannot have gone through
        assert not finished.wait(timeout=0.3)

    worker.join(timeout=5)
    assert finished.is_set()


def test_file_detail_follows_an_edit_under_the_same_fileid(tmp_path):
    root = make_repo(tmp_path)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/01_seed.sas"
    assert "set work.t3" in svc.file_detail(fileid)["code"]

    (root / "01_seed.sas").write_text("data work.t1; set work.t9; run;")
    ix.build()
    # the id is stable, so the same id must now answer with the NEW content
    detail = svc.file_detail(fileid)
    assert "set work.t9" in detail["code"]
    assert ("work.t9", "work.t1") in {(e["src"], e["dst"])
                                      for e in detail["block_edges"]}


def test_stored_engine_rows_follow_an_edit_under_the_same_fileid(tmp_path):
    root = make_repo(tmp_path)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/01_seed.sas"
    svc.excel_bytes([fileid])                      # first parse is persisted
    stored = lambda: {r[0] for r in store.con.execute(  # noqa: E731
        "SELECT source_canonical_name FROM lineage WHERE fileid = ?",
        [fileid]).fetchall()}
    assert "work.t3" in stored()

    (root / "01_seed.sas").write_text("data work.t1; set work.t9; run;")
    ix.build()
    svc.excel_bytes([fileid])
    # the id did not move, so the STORED rows have to be the new parse
    assert stored() == {"work.t9"}


def test_neighborhood_reaches_every_file_at_depth_2(svc):
    hood = svc.neighborhood(file=seed_fileid(svc), depth=2)
    assert {n["label"] for n in hood["nodes"]} == set(FIXTURE)


def test_neighborhood_edges_are_project_level_and_inferred(svc):
    hood = svc.neighborhood(file=seed_fileid(svc), depth=2)
    assert hood["edges"]
    assert {e["level"] for e in hood["edges"]} == {"project"}
    assert {e["provenance"] for e in hood["edges"]} == {"inferred"}
    assert all(e["src"] and e["dst"] and e["tables"] for e in hood["edges"])


def test_neighborhood_order_marks_the_loop(svc):
    hood = svc.neighborhood(file=seed_fileid(svc), depth=2)
    cyclic = {hood["order"][n["id"]]["file"] for n in hood["nodes"]
              if hood["order"][n["id"]]["cyclic"]}
    labels = {n["id"]: n["label"] for n in hood["nodes"]}
    assert {labels[f] for f in cyclic} == {
        "01_seed.sas", "02_mid.sas", "03_leaf.sas"}
    report = next(n for n in hood["nodes"] if n["label"] == "04_report.sas")
    assert not report["cyclic"] and report["score"] > 0


def test_story_is_data_flow_sentences(svc):
    hood = svc.neighborhood(file=seed_fileid(svc), depth=2)
    story = hood["story"]
    assert any("->" in line for line in story)
    assert any("downstream files are affected" in line for line in story)
    assert any("01_seed.sas -> work.t1 -> 02_mid.sas" in line for line in story)
    assert any("data feeds back into where it started" in line for line in story)
    # names, not raw fileids
    assert not any(seed_fileid(svc) in line for line in story)


def test_order_reasoning_reads_in_file_names(svc):
    hood = svc.neighborhood(file=seed_fileid(svc), depth=2)
    row = hood["order"][seed_fileid(svc)]
    assert row["label"] == "01_seed.sas"
    assert seed_fileid(svc) not in row["reasoning"]
    assert "work.t3 -> 01_seed.sas" in row["reasoning"]
    assert "break the flow 03_leaf.sas -> 01_seed.sas" in row["reasoning"]


def test_file_detail_block_edges_are_facts(svc):
    detail = svc.file_detail(seed_fileid(svc))
    assert detail["name"] == "01_seed.sas"
    assert "set work.t3" in detail["code"]
    assert detail["block_edges"]
    assert {e["provenance"] for e in detail["block_edges"]} == {"fact"}
    assert {e["level"] for e in detail["block_edges"]} == {"block"}
    assert ("work.t3", "work.t1") in {(e["src"], e["dst"])
                                      for e in detail["block_edges"]}
    assert all(e["provenance"] == "inferred" and e["level"] == "file"
               for e in detail["file_edges"])


def test_excel_has_three_sheets_and_the_provided_contract(svc):
    fileids = [n["id"] for n in
               svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    book = load_workbook(io.BytesIO(svc.excel_bytes(fileids)))
    assert book.sheetnames == ["provided", "customer_changes", "final_edges"]
    header = [c.value for c in book["provided"][1]]
    assert header[0] == "fileid"
    assert header[-2:] == ["edge_type", "provenance"]
    rows = list(book["provided"].iter_rows(min_row=2, values_only=True))
    assert rows
    by_type = {r[header.index("edge_type")]: r[header.index("provenance")]
               for r in rows}
    assert by_type["BLOCK_FLOW"] == "fact"
    assert by_type["PROJECT_FLOW"] == "inferred"


def test_add_edit_lands_in_customer_changes(svc):
    fileids = [n["id"] for n in
               svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    book = load_workbook(io.BytesIO(svc.excel_bytes(fileids)))
    assert list(book["customer_changes"].iter_rows(min_row=2,
                                                   values_only=True)) == []

    edit_id = svc.add_edit("confirm", "a.sas", "b.sas", "work.t1",
                           "project", "ravi", "ok")
    assert edit_id

    book = load_workbook(io.BytesIO(svc.excel_bytes(fileids)))
    sheet = book["customer_changes"]
    assert [c.value for c in sheet[1]] == service.CHANGE_COLUMNS
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 1
    assert rows[0][1:] == ("ravi", "confirm", "a.sas", "b.sas", "work.t1",
                           "project", "ok")


def test_add_edit_rejects_an_unknown_action(svc):
    with pytest.raises(ValueError):
        svc.add_edit("delete", "a.sas", "b.sas", "work.t1", "project",
                     "ravi", "")


def test_neighborhood_by_table(svc):
    hood = svc.neighborhood(table="work.t2", depth=1)
    assert {"02_mid.sas", "03_leaf.sas", "04_report.sas"} <= {
        n["label"] for n in hood["nodes"]}


def test_unknown_term_falls_back_to_a_live_scan(tmp_path):
    root = make_repo(tmp_path)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    # a file that landed after the index was built: the term is unknown, so
    # the seed has to come from a live scan of the tree
    (root / "05_late.sas").write_text("data work.late_mart; set work.t4; run;")
    assert ix.files_mentioning("work.late_mart") == []

    hood = service.Service(store, ix, [root]).neighborhood(
        table="work.late_mart", depth=1)
    labels = {n["label"] for n in hood["nodes"]}
    assert "05_late.sas" in labels and "04_report.sas" in labels
    assert any("04_report.sas -> work.t4 -> 05_late.sas" in line
               for line in hood["story"])


def test_neighborhood_is_deterministic(svc):
    a = svc.neighborhood(file=seed_fileid(svc), depth=2)
    b = svc.neighborhood(file=seed_fileid(svc), depth=2)
    assert a == b


# ---- UX round 2: directional depth (up/down) --------------------------------

# a straight chain with no loop, so upstream and downstream are distinct:
# 01_src -> work.t1 -> 02_mid -> work.t2 -> 03_dst
CHAIN = {
    "01_src.sas": "data work.t1; set rawlib.feed; run;",
    "02_mid.sas": "data work.t2; set work.t1; run;",
    "03_dst.sas": "data work.t3; set work.t2; run;",
}


@pytest.fixture
def chain_svc(tmp_path):
    root = tmp_path / "chain"
    root.mkdir()
    for name, code in CHAIN.items():
        (root / name).write_text(code)
    store = LineageStore(tmp_path / "chain.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    return service.Service(store, ix, [root])


def _labels_of(hood):
    return {n["label"] for n in hood["nodes"]}


def test_neighborhood_up_only_walks_against_the_flow(chain_svc):
    hood = chain_svc.neighborhood(file="chain/02_mid.sas", up=1, down=0)
    assert _labels_of(hood) == {"01_src.sas", "02_mid.sas"}


def test_neighborhood_down_only_walks_with_the_flow(chain_svc):
    hood = chain_svc.neighborhood(file="chain/02_mid.sas", up=0, down=1)
    assert _labels_of(hood) == {"02_mid.sas", "03_dst.sas"}


def test_neighborhood_depth_still_expands_both_ways(chain_svc):
    hood = chain_svc.neighborhood(file="chain/02_mid.sas", depth=1)
    assert _labels_of(hood) == set(CHAIN)


def test_neighborhood_nodes_carry_their_role(chain_svc):
    hood = chain_svc.neighborhood(file="chain/02_mid.sas", up=1, down=1)
    roles = {n["label"]: n["role"] for n in hood["nodes"]}
    assert roles == {"01_src.sas": "up", "02_mid.sas": "seed",
                     "03_dst.sas": "down"}


def test_neighborhood_names_its_seeds(chain_svc):
    hood = chain_svc.neighborhood(file="chain/02_mid.sas", up=1, down=1)
    assert hood["seeds"] == ["chain/02_mid.sas"]


def test_neighborhood_role_is_both_when_a_loop_reaches_a_file_both_ways(svc):
    # FIXTURE is a loop, so the seed's neighbors are up AND down at depth 2
    hood = svc.neighborhood(file=seed_fileid(svc), up=2, down=2)
    roles = {n["label"]: n["role"] for n in hood["nodes"]}
    assert roles["01_seed.sas"] == "seed"
    assert roles["02_mid.sas"] == "both"
    assert roles["03_leaf.sas"] == "both"


def test_neighborhood_table_seed_marks_its_files_as_seeds(chain_svc):
    hood = chain_svc.neighborhood(table="work.t1", up=0, down=0)
    roles = {n["label"]: n["role"] for n in hood["nodes"]}
    # both the writer and the reader of work.t1 mention it -> both are seeds
    assert roles == {"01_src.sas": "seed", "02_mid.sas": "seed"}


# ---- Task 13: edit lifecycle ------------------------------------------------

def test_edit_lifecycle_flags_stale_block_and_reconfirm_clears_it(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])                        # parse-on-demand ingest
    blocks_before = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())
    assert len(blocks_before) == 2
    block2_id = blocks_before[1]                      # "data work.t1; ..."

    edit_id = svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block",
                           "ravi", "looks right", block_id=block2_id)

    before = svc.final_edges([fileid])["rows"]
    assert any(e["block_id"] == block2_id and e["provenance"] == "human_gold"
              for e in before)

    # rewrite block 2 only (block 1 untouched) -> block 2's hash moves
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\n"
        "data work.t1; set work.t3; label t1='edited'; run;")
    ix.build()
    svc.excel_bytes([fileid])                         # re-ingest: hook fires

    # (a) the edit pointing at the vanished block is flagged
    flagged = store.con.execute(
        "SELECT requires_check FROM human_edits WHERE edit_id = ?",
        [edit_id]).fetchone()[0]
    assert flagged is True

    # (b) pending_checks() surfaces it
    pending = svc.pending_checks()
    assert edit_id in {row["edit_id"] for row in pending}
    assert pending[0]["block_id"] == block2_id

    # (c) final_edges no longer applies the flagged confirm
    mid = svc.final_edges([fileid])["rows"]
    assert not any(e["provenance"] == "human_gold" for e in mid)

    # (e) engine lineage/blocks rows for the file hold only the new block ids
    blocks_after = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())
    assert block2_id not in blocks_after
    assert len(blocks_after) == 2
    new_block2_id = blocks_after[1]

    # (d) re-confirming the same src -> dst -> table_name clears the flag
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi",
                "still right", block_id=new_block2_id)
    cleared = store.con.execute(
        "SELECT requires_check FROM human_edits WHERE edit_id = ?",
        [edit_id]).fetchone()[0]
    assert cleared is False
    assert svc.pending_checks() == []

    after = svc.final_edges([fileid])["rows"]
    assert any(e["block_id"] == new_block2_id and e["provenance"] == "human_gold"
              for e in after)


# ---- Task 10: final_edges paging/filtering + third Excel sheet -------------

def _first_block_id(svc, fileid) -> str:
    (block_id,) = sorted(r[0] for r in svc.store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())[:1]
    return block_id


def test_final_edges_reject_removes_the_matching_block(svc):
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])                       # ensure blocks are stored
    block_id = _first_block_id(svc, fileid)
    before = svc.final_edges([fileid])["rows"]
    assert any(r["block_id"] == block_id for r in before)

    svc.add_edit("reject", fileid, "work.t1", "work.t3", "block", "ravi",
                "wrong", block_id=block_id)
    after = svc.final_edges([fileid])["rows"]
    assert not any(r["block_id"] == block_id for r in after)


def test_final_edges_add_appears_as_human_gold(svc):
    svc.add_edit("add", "01_seed.sas", "05_extra.sas", "work.extra",
                "project", "ravi", "new flow")
    rows = svc.final_edges()["rows"]
    added = [r for r in rows if r["dst"] == "05_extra.sas"]
    assert added and added[0]["provenance"] == "human_gold"
    assert added[0]["src"] == "01_seed.sas"
    assert added[0]["tables"] == ["work.extra"]
    assert added[0]["level"] == "project"


def test_final_edges_confirm_upgrades_provenance(svc):
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    block_id = _first_block_id(svc, fileid)
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi",
                "ok", block_id=block_id)
    rows = svc.final_edges([fileid])["rows"]
    matched = [r for r in rows if r["block_id"] == block_id]
    assert matched
    assert all(r["provenance"] == "human_gold" for r in matched)


def test_final_edges_filter_and_level_and_paging(svc):
    fileids = [n["id"] for n in
              svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    all_rows = svc.final_edges(fileids)["total"]
    assert all_rows > 1

    filtered = svc.final_edges(fileids, filter_text="t3")
    assert filtered["total"] >= 1
    assert all("t3" in r["src"].lower() or "t3" in r["dst"].lower()
              or any("t3" in t.lower() for t in r["tables"])
              for r in filtered["rows"])

    by_level = svc.final_edges(fileids, level="block")
    assert by_level["rows"]
    assert all(r["level"] == "block" for r in by_level["rows"])

    page = svc.final_edges(fileids, limit=1)
    assert page["total"] == all_rows
    assert len(page["rows"]) == 1

    beyond = svc.final_edges(fileids, offset=all_rows + 100)
    assert beyond["rows"] == []
    assert beyond["total"] == all_rows


def test_excel_has_three_sheets_with_hyperlinked_final_edges(svc):
    fileids = [n["id"] for n in
              svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    book = load_workbook(io.BytesIO(svc.excel_bytes(fileids)))
    assert book.sheetnames == ["provided", "customer_changes", "final_edges"]
    header = [c.value for c in book["final_edges"][1]]
    assert header == service.FINAL_EDGE_COLUMNS
    rows = list(book["final_edges"].iter_rows(min_row=2, values_only=True))
    assert rows
    assert all(r[-1].startswith("=HYPERLINK(") for r in rows)


def test_excel_final_edges_overflows_past_the_row_cap(svc, monkeypatch):
    monkeypatch.setattr(service, "MAX_XLSX_ROWS", 1)
    fileids = [n["id"] for n in
              svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    book = load_workbook(io.BytesIO(svc.excel_bytes(fileids)))
    rows = list(book["final_edges"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2                            # 1 real row + overflow
    assert rows[-1][0] == "TRUNCATED — scope the download"
    # openpyxl round-trips an appended "" as an empty (None) cell
    assert rows[-1][1:] == (None, None, None, None, None)


# ---- Task 14: Arrow IPC edge views (engine | human | final) -----------------

def _arrow_table(payload: bytes):
    """The IPC stream the browser receives, decoded back into a table."""
    import pyarrow as pa

    return pa.ipc.open_stream(pa.BufferReader(payload)).read_all()


def test_edges_arrow_columns_are_the_grid_contract(svc):
    fileid = seed_fileid(svc)
    table = _arrow_table(svc.edges_arrow([fileid], view="engine"))
    assert table.schema.names == service.EDGE_ARROW_COLUMNS
    assert table.num_rows > 0


def test_edges_arrow_engine_view_is_the_tools_own_rows(svc):
    """engine = the latest parse, read-only: a customer's confirm never
    changes it (that is what the final view is for)."""
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    block_id = _first_block_id(svc, fileid)
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi",
                 "ok", block_id=block_id)

    engine = _arrow_table(svc.edges_arrow([fileid], view="engine")).to_pydict()
    assert set(engine["provenance"]) <= {"fact", "inferred"}
    assert all(flag is False for flag in engine["requires_check"])

    final = _arrow_table(svc.edges_arrow([fileid], view="final")).to_pydict()
    assert "human_gold" in final["provenance"]


def test_edges_arrow_final_view_matches_the_paged_json(svc):
    fileids = [n["id"] for n in
               svc.neighborhood(file=seed_fileid(svc), depth=2)["nodes"]]
    table = _arrow_table(svc.edges_arrow(fileids, view="final"))
    assert table.num_rows == svc.final_edges(fileids)["total"]


def test_edges_arrow_human_view_carries_requires_check(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])
    block2_id = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())[1]
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi",
                 "looks right", block_id=block2_id)

    human = _arrow_table(svc.edges_arrow([fileid], view="human")).to_pydict()
    assert human["src"] == [fileid]
    assert human["dst"] == ["work.t1"]
    assert human["provenance"] == ["human_gold"]
    assert human["requires_check"] == [False]

    # the block this edit was made against changes -> the edit is queued for
    # review, and the human view is where that queue is visible
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\n"
        "data work.t1; set work.t3; label t1='edited'; run;")
    ix.build()
    svc.excel_bytes([fileid])

    flagged = _arrow_table(svc.edges_arrow([fileid], view="human")).to_pydict()
    assert flagged["requires_check"] == [True]


def test_edges_arrow_rejects_an_unknown_view(svc):
    with pytest.raises(ValueError):
        svc.edges_arrow([seed_fileid(svc)], view="excel")


def test_every_arrow_view_shares_the_same_column_contract(svc):
    """engine | human | final all hand the grid the SAME columns, `fileid`
    first — the grid opens an edge's SAS code by that column, and a block id
    alone cannot name its file (the same id can exist in several files)."""
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    block_id = _first_block_id(svc, fileid)
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi",
                 "ok", block_id=block_id)

    for view in service.ARROW_VIEWS:
        table = svc.edges_arrow_table([fileid], view=view)
        assert table.schema.names == service.EDGE_ARROW_COLUMNS, view
        assert table.num_rows > 0, view
        # never NULL: the browser uses it as a route key
        assert all(f is not None
                   for f in table.column("fileid").to_pylist()), view


def test_engine_and_final_rows_carry_the_file_the_edge_lives_in(svc):
    fileid = seed_fileid(svc)
    for view in ("engine", "final"):
        table = svc.edges_arrow_table([fileid], view=view)
        assert set(table.column("fileid").to_pylist()) == {fileid}


# ---- line_start: where a block starts in its file's source ------------------

STEPS_SAS = (
    "/* a header comment\n"                      # 1
    "   that spans two lines */\n"               # 2
    "\n"                                         # 3
    "data work.a;\n"                             # 4  <- block 1
    "  set work.x;\n"                            # 5
    "run;\n"                                     # 6
    "\n"                                         # 7
    "/* data work.ghost; set work.y; run; */\n"  # 8  (commented out)
    "proc sort data=work.a out=work.b;\n"        # 9  <- block 2
    "  by id;\n"                                 # 10
    "run;\n"                                     # 11
    "\n"                                         # 12
    "data work.c;\n"                             # 13 <- block 3
    "  set work.b;\n"                            # 14
    "run;\n"                                     # 15
)

ANNOTATED_SAS = (
    "/*BLOCKID b_1_deadbeef, Data Step, Lines: 3;*/\n"   # 1
    "data work.a;\n"                                     # 2
    "  set work.x;\n"                                    # 3
    "run;\n"                                             # 4
    "/*ENDBLOCKID b_1_deadbeef, Data Step;*/\n"          # 5
)


def test_block_line_starts_counts_step_headers_past_blanks_and_comments():
    """Positional fallback: the i-th block takes the i-th `data`/`proc`
    header line — blank lines and commented-out steps never count."""
    assert service.block_line_starts(STEPS_SAS, ["b1", "b2", "b3"]) == [
        4, 9, 13]


def test_block_line_starts_prefers_an_injected_blockid_annotation():
    """A block id that appears literally in the source resolves to ITS OWN
    annotation line; ids with no annotation still fall back positionally."""
    assert service.block_line_starts(
        ANNOTATED_SAS, ["b_1_deadbeef"]) == [1]
    # id 2 has no annotation and no second header -> the last header found
    assert service.block_line_starts(
        ANNOTATED_SAS, ["b_1_deadbeef", "b_2_nothere"]) == [1, 2]


def test_block_line_starts_never_raises_and_always_returns_line_one_at_worst():
    assert service.block_line_starts("", ["b1"]) == [1]
    assert service.block_line_starts("* just a note;\n", ["b1", "b2"]) == [1, 1]
    assert service.block_line_starts(STEPS_SAS, []) == []


def test_file_detail_blocks_carry_a_one_based_line_start(svc):
    detail = svc.file_detail(seed_fileid(svc))
    assert detail["blocks"]
    for block in detail["blocks"]:
        assert isinstance(block["line_start"], int)
        assert block["line_start"] >= 1
    # the fixture file is one single-line step
    assert detail["blocks"][0]["line_start"] == 1


def test_file_detail_line_starts_point_at_each_step(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "steps.sas").write_text(STEPS_SAS)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    detail = svc.file_detail("repo/steps.sas")
    assert [b["line_start"] for b in detail["blocks"]] == [4, 9, 13]


# ---- `correct` edits + the file an edit belongs to --------------------------

def test_correct_is_one_of_the_recorded_actions():
    assert service.ACTIONS == ("add", "confirm", "correct", "reject")


def test_correct_adds_the_right_flow_without_dropping_the_engine_row(svc):
    """`correct` reads as "the engine row was wrong, here is the right
    flow": the corrected flow lands as human_gold and the engine's own row
    stays — removing a row is what `reject` is for."""
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    block_id = _first_block_id(svc, fileid)
    assert any(r["block_id"] == block_id
               for r in svc.final_edges([fileid])["rows"])

    svc.add_edit("correct", fileid, "work.t1", "work.right", "block", "ravi",
                 "reads work.right, not work.t3", block_id=block_id,
                 fileid=fileid)

    rows = svc.final_edges([fileid])["rows"]
    corrected = [r for r in rows if r["tables"] == ["work.right"]]
    assert corrected and corrected[0]["provenance"] == "human_gold"
    assert corrected[0]["fileid"] == fileid
    assert any(r["block_id"] == block_id and r["provenance"] != "human_gold"
               for r in rows)


def test_correct_shows_in_the_human_view_carrying_its_fileid(svc):
    fileid = seed_fileid(svc)
    svc.add_edit("correct", fileid, "work.t1", "work.right", "block", "ravi",
                 "", block_id="b_1_abcdef01", fileid=fileid)
    human = svc.edges_arrow_table([fileid], view="human").to_pydict()
    assert human["fileid"] == [fileid]
    assert human["table_name"] == ["work.right"]
    assert human["provenance"] == ["human_gold"]


def test_an_edit_without_a_fileid_still_has_the_column(svc):
    """A file/project-level edit names no single file — the human view still
    hands the grid a `fileid` cell, empty rather than NULL."""
    svc.add_edit("add", "01_seed.sas", "05_extra.sas", "work.extra",
                 "project", "ravi", "new flow")
    human = svc.edges_arrow_table([seed_fileid(svc)],
                                  view="human").to_pydict()
    assert human["fileid"] == [""]


def test_pending_checks_reports_the_file_the_edit_was_made_in(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])
    block2_id = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())[1]
    svc.add_edit("correct", fileid, "work.t1", "work.right", "block", "ravi",
                 "", block_id=block2_id, fileid=fileid)

    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\n"
        "data work.t1; set work.t3; label t1='edited'; run;")
    ix.build()
    svc.excel_bytes([fileid])

    (pending,) = svc.pending_checks()
    assert pending["fileid"] == fileid
    assert pending["action"] == "correct"


# ---- T1: dismissed column — reject retires stale edits from the queue ------

def test_rejecting_a_stale_edit_dismisses_it_from_the_queue(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])

    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])                       # parse-on-demand ingest
    blocks = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())
    edit_id = svc.add_edit("add", fileid, "work.t9", "work.extra", "block",
                           "ravi", "extra flow", block_id=blocks[1])
    # make the pinned block vanish -> the edit gets flagged
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\n"
        "data work.t1; set work.t3; label t1='edited'; run;")
    ix.build()
    svc.excel_bytes([fileid])
    assert edit_id in {r["edit_id"] for r in svc.pending_checks()}

    # discard: reject the same flow -> queue empties...
    svc.add_edit("reject", fileid, "work.t9", "work.extra", "block",
                 "ravi", "no longer true")
    assert svc.pending_checks() == []

    # ...and the dismissed edit did NOT re-enter the merge
    dismissed, flagged = store.con.execute(
        "SELECT dismissed, requires_check FROM human_edits "
        "WHERE edit_id = ?", [edit_id]).fetchone()
    assert dismissed is True and flagged is True
    rows = svc.final_edges([fileid])["rows"]
    assert not any("work.extra" in r["tables"]
                   and r["provenance"] == "human_gold" for r in rows)


def test_confirming_still_clears_the_flag_and_pending(tmp_path):
    # guard: the existing confirm path must keep working with the new column
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])
    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])
    blocks = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block",
                 "ravi", "", block_id=blocks[1])
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\n"
        "data work.t1; set work.t3; label t1='x'; run;")
    ix.build()
    svc.excel_bytes([fileid])
    assert len(svc.pending_checks()) == 1
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block", "ravi", "")
    assert svc.pending_checks() == []


# ---- code save: shadow-git timeline + edge delta ---------------------------

def _save_fixture(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "edit.sas").write_text(
        "data work.a; set work.x; run;\ndata work.t1; set work.t3; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    from history import EstateHistory
    svc = service.Service(store, ix, [root],
                          history=EstateHistory(tmp_path / "hist"))
    return root, store, svc


def test_save_code_writes_disk_commits_and_reports_edge_delta(tmp_path):
    root, store, svc = _save_fixture(tmp_path)
    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])                       # ingest old bytes

    new_code = ("data work.a; set work.x; run;\n"
                "data work.t1; set work.t9; run;")  # t3 -> t9
    result = svc.save_code(fileid, new_code, "ravi", "switch to t9 feed")

    assert (root / "edit.sas").read_text() == new_code
    assert {"src": "work.t9", "dst": "work.t1", "table_name": "work.t9",
            "level": "block"} in result["added"]
    assert {"src": "work.t3", "dst": "work.t1", "table_name": "work.t3",
            "level": "block"} in result["removed"]
    assert result["commit"]

    hist = svc.history.history(fileid)
    assert [h["editor"] for h in hist] == ["ravi", "engine"]
    assert hist[1]["reason"] == "original import"
    old = svc.history.at(fileid, hist[1]["commit"])
    assert "set work.t3" in old["code"]


def test_save_code_flags_edits_pinned_to_changed_blocks(tmp_path):
    root, store, svc = _save_fixture(tmp_path)
    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])
    blocks = sorted(r[0] for r in store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = ?", [fileid]).fetchall())
    svc.add_edit("confirm", fileid, "work.t1", "work.t3", "block",
                 "ravi", "", block_id=blocks[1])

    svc.save_code(fileid,
                  "data work.a; set work.x; run;\n"
                  "data work.t1; set work.t3; label t1='x'; run;",
                  "ravi", "cosmetic label")

    assert len(svc.pending_checks()) == 1


def test_save_code_requires_editor_and_reason(tmp_path):
    _root, _store, svc = _save_fixture(tmp_path)
    fileid = "repo/edit.sas"
    svc.excel_bytes([fileid])
    with pytest.raises(ValueError):
        svc.save_code(fileid, "data a; run;", "", "reason")
    with pytest.raises(ValueError):
        svc.save_code(fileid, "data a; run;", "ravi", "  ")


# ---- store-backed scanner (sources) ----------------------------------------

def test_scan_sources_term_matches_python_scanner(svc):
    for fid in [r["id"] for r in svc.list_files()]:
        svc._ensure_ingested([fid])
    hits = svc.scan_sources_term("WORK.T1")               # case-insensitive
    assert hits == ["repo/01_seed.sas", "repo/02_mid.sas"]
    assert svc.scan_sources_term("no.such_table") == []


def test_sources_ready_only_after_matching_sync_meta(svc):
    assert svc.sources_ready() is False
    svc.store.con.execute(
        "INSERT INTO sync_meta VALUES ('synced_roots', ?)",
        [svc.roots_signature()])
    assert svc.sources_ready() is True


def test_sources_ready_false_for_other_roots(svc):
    svc.store.con.execute(
        "INSERT INTO sync_meta VALUES ('synced_roots', '/other/place')")
    assert svc.sources_ready() is False


# ---- sync_sources (tree-wide hash-gated refresh) ----------------------------

def test_sync_sources_ingests_everything_and_marks_ready(svc):
    report = svc.sync_sources()
    assert report["files_on_disk"] == 4
    assert report["ingested"] == 4
    assert report["removed"] == 0
    assert svc.sources_ready() is True
    assert svc.store.con.execute(
        "SELECT count(*) FROM sources").fetchone()[0] == 4


def test_sync_sources_skips_unchanged_reingests_changed(svc, tmp_path):
    svc.sync_sources()
    (tmp_path / "repo" / "01_seed.sas").write_text(
        "data work.t1; set work.t9; run;")
    report = svc.sync_sources()
    assert report["ingested"] == 1                    # only the changed file
    src = svc.store.con.execute(
        "SELECT source FROM sources WHERE fileid = 'repo/01_seed.sas'"
    ).fetchone()[0]
    assert "work.t9" in src


def test_sync_sources_removes_vanished_files(svc, tmp_path):
    svc.sync_sources()
    (tmp_path / "repo" / "04_report.sas").unlink()
    report = svc.sync_sources()
    assert report["removed"] == 1
    assert svc.store.con.execute(
        "SELECT count(*) FROM files WHERE fileid = 'repo/04_report.sas'"
    ).fetchone()[0] == 0
    assert svc.store.con.execute(
        "SELECT count(*) FROM sources WHERE fileid = 'repo/04_report.sas'"
    ).fetchone()[0] == 0


# ---- freshness in the edge views --------------------------------------------

def test_arrow_final_view_carries_freshness_and_red_ghosts(svc, tmp_path):
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])                        # ingest baseline
    (tmp_path / "repo" / "01_seed.sas").write_text(
        "data work.t1; set work.t9; run;")           # t3 -> t1 becomes t9 -> t1
    svc.excel_bytes([fileid])                        # re-ingest
    table = svc.edges_arrow_table([fileid], view="final")
    cols = {name: table.column(name).to_pylist()
            for name in table.schema.names}
    assert "freshness" in cols
    rows = list(zip(cols["src"], cols["dst"], cols["freshness"]))
    # the vanished engine flow rides along as a red ghost row
    assert ("work.t3", "work.t1", "red") in rows
    # red sorts first
    assert cols["freshness"][0] == "red"


def test_yellow_edit_keeps_confirm_red_edit_flags(svc, tmp_path):
    """Yellow keeps confirm (approved decision); red joins the queue."""
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    block_id = _first_block_id(svc, fileid)
    edit_id = svc.add_edit("confirm", "work.t3", "work.t1", "work.t3",
                           "block", "ravi", "", block_id=block_id,
                           fileid=fileid)
    # cosmetic change: same flow, new block hash -> yellow, confirm kept
    (tmp_path / "repo" / "01_seed.sas").write_text(
        "data work.t1; set work.t3; label t1='x'; run;")
    svc.excel_bytes([fileid])
    flagged, fresh = svc.store.con.execute(
        "SELECT requires_check, freshness FROM human_edits "
        "WHERE edit_id = ?", [edit_id]).fetchone()
    assert flagged is False and fresh == "yellow"
    # flow removed entirely -> red, requires_check raised
    (tmp_path / "repo" / "01_seed.sas").write_text(
        "data work.t1; set work.t9; run;")
    svc.excel_bytes([fileid])
    flagged, fresh = svc.store.con.execute(
        "SELECT requires_check, freshness FROM human_edits "
        "WHERE edit_id = ?", [edit_id]).fetchone()
    assert flagged is True and fresh == "red"


def test_edges_json_rows_carry_freshness(svc):
    fileid = seed_fileid(svc)
    rows = svc.final_edges([fileid])["rows"]
    assert rows and all(r["freshness"] == "green" for r in rows)


def test_file_detail_carries_line_ranges_and_verify_fields(svc):
    fileid = seed_fileid(svc)
    detail = svc.file_detail(fileid)
    assert "macro_calls" in detail
    assert "includes" in detail and "missing_includes" in detail
    total_lines = len(detail["code"].split("\n"))
    for block in detail["blocks"]:
        assert block["line_end"] >= block["line_start"]
        assert block["line_end"] <= total_lines
        assert block["occurrences"]
    # freshly parsed: every edge is green
    for edge in detail["block_edges"] + detail["file_edges"]:
        assert edge["freshness"] == "green"


def test_any_edit_on_a_lost_flow_clears_its_ghost(svc, tmp_path):
    fileid = seed_fileid(svc)
    svc.excel_bytes([fileid])
    (tmp_path / "repo" / "01_seed.sas").write_text(
        "data work.t1; set work.t9; run;")
    svc.excel_bytes([fileid])
    assert svc.store.con.execute(
        "SELECT count(*) FROM lost_flows").fetchone()[0] == 1
    svc.add_edit("reject", "work.t3", "work.t1", "work.t3", "block",
                 "ravi", "gone for real")
    assert svc.store.con.execute(
        "SELECT count(*) FROM lost_flows").fetchone()[0] == 0


# ---- Task V2-3: block_links + (fileid, block_id)-pinned edits --------------

def test_block_links_returns_block_level_cross_file_rows(tmp_path):
    root = tmp_path / "repo"
    root.mkdir()
    (root / "w.sas").write_text("data work.t; set raw.x; run;")
    (root / "r.sas").write_text("data work.out; set work.t; run;")
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])
    links = svc.block_links(["repo/w.sas", "repo/r.sas"])
    assert len(links) == 1
    l = links[0]
    assert l["src_file"] == "repo/w.sas" and l["dst_file"] == "repo/r.sas"
    assert l["table"] == "work.t"
    assert l["src_ref"].startswith(l["src_block"] + ":t_")


def test_reject_pinned_to_a_block_only_hits_that_file(tmp_path):
    # two files with IDENTICAL content -> identical content-hashed block ids
    root = tmp_path / "repo"
    root.mkdir()
    code = "data work.t; set raw.x; run;"
    (root / "a.sas").write_text(code)
    (root / "b.sas").write_text(code)
    store = LineageStore(tmp_path / "db.duckdb")
    ix = Indexer(store, [root])
    ix.build()
    svc = service.Service(store, ix, [root])
    svc.excel_bytes(["repo/a.sas", "repo/b.sas"])   # parse-on-demand ingest
    block = store.con.execute(
        "SELECT block_id FROM blocks WHERE fileid = 'repo/a.sas'").fetchone()[0]
    svc.add_edit("reject", "raw.x", "work.t", "work.t", "block",
                 "ravi", "wrong in a only", block_id=block, fileid="repo/a.sas")
    rows = svc.final_edges(["repo/a.sas", "repo/b.sas"])["rows"]
    block_rows = [r for r in rows if r["level"] == "block"]
    assert all(r["fileid"] != "repo/a.sas" for r in block_rows)      # a: dropped
    assert any(r["fileid"] == "repo/b.sas" for r in block_rows)      # b: SURVIVES
